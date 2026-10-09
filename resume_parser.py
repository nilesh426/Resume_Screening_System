"""
resume_parser.py - Extract text from PDF and DOCX resumes.
Beginner-friendly: no files are saved to disk, everything is in-memory.
Never executes uploaded files - only reads text.
"""

import io

ALLOWED_EXTENSIONS = {"pdf", "docx"}


def allowed_file(filename):
    """Check if filename has an allowed extension."""
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_text_from_pdf(file_bytes):
    """Extract text from PDF bytes using pypdf.

    Returns (text, error). If text cannot be extracted (e.g. scanned PDF),
    text will be an empty string and error will explain why.
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        # Fallback for older requirement name PyPDF2
        from PyPDF2 import PdfReader

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        parts = []
        for page in reader.pages:
            try:
                t = page.extract_text()
            except Exception:
                t = None
            if t:
                parts.append(t)
        text = "\n".join(parts).strip()
        if not text:
            return "", "No extractable text found (possibly a scanned image PDF)."
        return text, None
    except Exception as e:
        return "", f"Could not read PDF file: {e}"


def extract_text_from_docx(file_bytes):
    """Extract text from DOCX bytes using python-docx."""
    try:
        import docx
    except ImportError:
        return "", "python-docx is not installed. Run: pip install python-docx"

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        # Also read tables (skills tables are common in resumes)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        parts.append(cell.text.strip())
        text = "\n".join(parts).strip()
        if not text:
            return "", "DOCX file contains no readable text."
        return text, None
    except Exception as e:
        return "", f"Could not read DOCX file: {e}"


def extract_text(filename, file_bytes):
    """Dispatch to the correct extractor based on extension.

    Returns (text, error).
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext == "docx":
        return extract_text_from_docx(file_bytes)
    else:
        return "", f"Unsupported file type '.{ext}'. Only PDF and DOCX are allowed."


def unique_filename(filename, seen):
    """Handle duplicate filenames: resume.pdf -> resume (1).pdf etc."""
    if filename not in seen:
        return filename
    name, dot, ext = filename.rpartition(".")
    if not dot:
        name, ext, dot = filename, "", ""
    i = 1
    while True:
        candidate = f"{name} ({i}).{ext}" if dot else f"{name} ({i})"
        if candidate not in seen:
            return candidate
        i += 1

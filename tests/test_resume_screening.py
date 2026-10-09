"""Tests for the Resume Screening System (no redesign, no new features).

Covers: single + multiple uploads, PDF/DOCX extraction, JD input and
match-score calculation, descending sort order, and graceful handling of
empty JDs, invalid types, unreadable PDFs, empty text and missing uploads.

Run from the project root:
    python -m pytest -q
"""

import io

import pytest

import app as app_mod
from nlp_processor import (
    _fallback_tokens,
    clean_for_tfidf,
    extract_keywords,
    screen_resumes,
)
from resume_parser import allowed_file, extract_text, unique_filename

SAMPLE_JD = (
    "Hiring Python developer with Flask web framework, SQL databases, "
    "REST APIs, pandas data analysis, machine learning basics, Git teamwork, "
    "debugging, testing, documentation and collaboration skills required."
)

STRONG_RESUME = (
    "Python developer with Flask web framework experience, SQL databases, "
    "REST APIs, pandas data analysis, Git teamwork, debugging and testing."
)
WEAK_RESUME = (
    "Sales marketing retail executive focused on customer service, "
    "store operations and hospitality management experience."
)


# --------------------------------------------------------------------------
# In-memory file builders (no fixture files, no extra dependencies)
# --------------------------------------------------------------------------

def make_docx_bytes(paragraphs=None, table_cells=None):
    """Build a minimal DOCX file in memory."""
    from docx import Document

    doc = Document()
    for p in paragraphs or []:
        doc.add_paragraph(p)
    if table_cells:
        table = doc.add_table(rows=1, cols=len(table_cells))
        for i, val in enumerate(table_cells):
            table.rows[0].cells[i].text = val
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def make_pdf_bytes(lines):
    """Build a minimal one-page text PDF in memory (hand-crafted)."""
    def esc(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    ops = ("BT /F1 12 Tf 72 720 Td 14 TL "
           + " T* ".join("(%s) Tj" % esc(line) for line in lines) + " ET")
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        ("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"),
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        "<< /Length %d >>\nstream\n%s\nendstream" % (len(ops), ops),
    ]
    out = [b"%PDF-1.4"]
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(sum(len(chunk) + 1 for chunk in out))
        out.append(("%d 0 obj\n%s\nendobj" % (i, body)).encode("latin-1"))
    xref_pos = sum(len(chunk) + 1 for chunk in out)
    out.append(("xref\n0 %d\n0000000000 65535 f " % (len(objs) + 1)).encode())
    for off in offsets:
        out.append(("%010d 00000 n " % off).encode())
    out.append(("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF"
                % (len(objs) + 1, xref_pos)).encode("latin-1"))
    return b"\n".join(out)


def make_blank_pdf_bytes():
    """A valid PDF page with no text layer (simulates a scanned PDF)."""
    from pypdf import PdfWriter

    buf = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    writer.write(buf)
    return buf.getvalue()


@pytest.fixture()
def client():
    app_mod.app.config["TESTING"] = True
    return app_mod.app.test_client()


def post_screen(client, jd, files):
    return client.post("/screen",
                       data={"job_description": jd, "resumes": files},
                       content_type="multipart/form-data")


# --------------------------------------------------------------------------
# 1. File-type validation and filename handling
# --------------------------------------------------------------------------

class TestAllowedFiles:
    @pytest.mark.parametrize("name", ["cv.pdf", "cv.docx", "CV.PDF",
                                      "my resume.DoCx"])
    def test_allowed(self, name):
        assert allowed_file(name) is True

    @pytest.mark.parametrize("name", ["cv.txt", "cv.exe", "cv", "cv.pdf.exe",
                                      "", "resume.zip"])
    def test_rejected(self, name):
        assert allowed_file(name) is False


class TestUniqueFilenames:
    def test_first_use_unchanged(self):
        assert unique_filename("cv.pdf", set()) == "cv.pdf"

    def test_duplicate_gets_suffix(self):
        assert unique_filename("cv.pdf", {"cv.pdf"}) == "cv (1).pdf"

    def test_suffix_chain(self):
        seen = {"cv.pdf", "cv (1).pdf"}
        assert unique_filename("cv.pdf", seen) == "cv (2).pdf"


# --------------------------------------------------------------------------
# 2. PDF and DOCX text extraction
# --------------------------------------------------------------------------

class TestExtraction:
    def test_docx_paragraphs(self):
        text, err = extract_text("cv.docx",
                                 make_docx_bytes(["Python developer",
                                                  "Flask experience"]))
        assert err is None
        assert "Python developer" in text
        assert "Flask experience" in text

    def test_docx_tables(self):
        text, err = extract_text("cv.docx", make_docx_bytes(
            ["Data analyst"], table_cells=["Python", "SQL"]))
        assert err is None
        assert "Python" in text and "SQL" in text

    def test_docx_empty(self):
        text, err = extract_text("empty.docx", make_docx_bytes([]))
        assert text == "" and err is not None

    def test_docx_corrupt(self):
        text, err = extract_text("bad.docx", b"this is not a docx file")
        assert text == "" and err is not None

    def test_pdf_text_multiline(self):
        text, err = extract_text("cv.pdf", make_pdf_bytes(
            ["Python Flask developer", "SQL and REST APIs"]))
        assert err is None
        assert "Python Flask developer" in text
        assert "SQL and REST APIs" in text

    def test_pdf_blank_page_reports_no_text(self):
        text, err = extract_text("scan.pdf", make_blank_pdf_bytes())
        assert text == ""
        assert err is not None and "extractable text" in err

    def test_pdf_corrupt_bytes(self):
        text, err = extract_text("bad.pdf", b"this is not a pdf file")
        assert text == "" and err is not None

    def test_unsupported_extension(self):
        text, err = extract_text("notes.txt", b"hello")
        assert text == "" and "Unsupported" in err


# --------------------------------------------------------------------------
# 3. NLP preprocessing, scoring and ranking
# --------------------------------------------------------------------------

class TestNlp:
    def test_fallback_drops_stopwords_and_short_tokens(self):
        tokens = _fallback_tokens("The Python developer and an AI go to SQL")
        assert "python" in tokens and "developer" in tokens
        assert "the" not in tokens and "and" not in tokens
        assert "ai" not in tokens  # too short

    def test_clean_empty_input(self):
        assert clean_for_tfidf("", None) == ""
        assert clean_for_tfidf("   ", None) == ""

    def test_extract_keywords_empty(self):
        assert extract_keywords("", None) == []

    def test_strong_resume_outranks_weak(self):
        ranked, _ = screen_resumes(SAMPLE_JD, {"strong.docx": STRONG_RESUME,
                                               "weak.docx": WEAK_RESUME})
        assert [r["filename"] for r in ranked] == ["strong.docx", "weak.docx"]
        assert ranked[0]["score"] > ranked[1]["score"]
        for r in ranked:
            assert 0.0 <= r["score"] <= 100.0

    def test_ranks_sequential_descending(self):
        texts = {"a.docx": STRONG_RESUME,
                 "b.docx": "Flask SQL developer with Git",
                 "c.docx": WEAK_RESUME}
        ranked, _ = screen_resumes(SAMPLE_JD, texts)
        scores = [r["score"] for r in ranked]
        assert scores == sorted(scores, reverse=True)
        assert [r["rank"] for r in ranked] == [1, 2, 3]

    def test_matched_and_missing_keywords(self):
        ranked, job_kw = screen_resumes(
            "Python Flask developer with kubernetes experience required",
            {"cv.docx": "Python Flask developer with SQL databases"})
        assert "python" in ranked[0]["matched_keywords"]
        assert "flask" in ranked[0]["matched_keywords"]
        assert "python" not in ranked[0]["missing_keywords"]
        assert "kubernetes" in ranked[0]["missing_keywords"]
        assert len(ranked[0]["missing_keywords"]) <= 20
        assert "python" in job_kw and "kubernetes" in job_kw

    def test_stopwords_only_jd_does_not_crash(self):
        ranked, _ = screen_resumes("the and or the and of of",
                                   {"a.docx": STRONG_RESUME})
        assert ranked[0]["score"] == 0.0

    def test_empty_resume_dict(self):
        ranked, job_kw = screen_resumes("Python developer needed", {})
        assert ranked == [] and "python" in job_kw

    def test_identical_resumes_tie(self):
        ranked, _ = screen_resumes(SAMPLE_JD, {"a.docx": STRONG_RESUME,
                                               "b.docx": STRONG_RESUME})
        assert ranked[0]["score"] == ranked[1]["score"]


# --------------------------------------------------------------------------
# 4. Flask routes: uploads, validation, results
# --------------------------------------------------------------------------

class TestRoutes:
    def test_index_renders_form(self, client):
        html = client.get("/").data.decode()
        assert 'name="resumes"' in html
        assert 'name="job_description"' in html
        assert "multiple" in html

    def test_single_docx_upload(self, client):
        html = post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "solo.docx")
        ]).data.decode()
        assert "Ranked Candidates (1)" in html
        assert "solo.docx" in html

    def test_multiple_uploads_all_ranked_descending(self, client):
        html = post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([WEAK_RESUME])), "weak.docx"),
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "strong.docx"),
            (io.BytesIO(make_pdf_bytes(["Flask SQL developer"])),
             "mid.pdf"),
        ]).data.decode()
        assert "Ranked Candidates (3)" in html
        assert (html.index("strong.docx") < html.index("mid.pdf")
                < html.index("weak.docx"))

    def test_pdf_and_docx_mix(self, client):
        html = post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_pdf_bytes(["Python Flask SQL developer"])),
             "a.pdf"),
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "b.docx"),
        ]).data.decode()
        assert "Ranked Candidates (2)" in html

    def test_empty_jd_rejected(self, client):
        resp = client.post("/screen",
                           data={"job_description": "   ",
                                 "resumes": [(io.BytesIO(
                                     make_docx_bytes([STRONG_RESUME])),
                                     "a.docx")]},
                           content_type="multipart/form-data",
                           follow_redirects=True)
        assert "job description" in resp.data.decode().lower()

    def test_missing_upload_rejected(self, client):
        html = client.post("/screen", data={"job_description": SAMPLE_JD},
                           content_type="multipart/form-data",
                           follow_redirects=True).data.decode()
        assert "upload at least one" in html.lower()

    def test_invalid_type_only(self, client):
        resp = client.post("/screen",
                           data={"job_description": SAMPLE_JD,
                                 "resumes": [(io.BytesIO(b"MZ junk"),
                                              "run.exe")]},
                           content_type="multipart/form-data",
                           follow_redirects=True)
        html = resp.data.decode()
        assert "No resumes could be processed" in html

    def test_one_invalid_file_does_not_crash_batch(self, client):
        html = post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "good.docx"),
            (io.BytesIO(b"MZ junk"), "bad.exe"),
            (io.BytesIO(make_blank_pdf_bytes()), "scan.pdf"),
        ]).data.decode()
        assert "Ranked Candidates (1)" in html
        assert "good.docx" in html
        assert "Skipped" in html

    def test_unreadable_pdf_skipped_with_warning(self, client):
        html = post_screen(client, SAMPLE_JD, [
            (io.BytesIO(b"not a pdf"), "broken.pdf"),
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "ok.docx"),
        ]).data.decode()
        assert "Ranked Candidates (1)" in html
        assert "broken.pdf" in html

    def test_empty_docx_skipped(self, client):
        html = post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([])), "empty.docx"),
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "ok.docx"),
        ]).data.decode()
        assert "Ranked Candidates (1)" in html

    def test_duplicate_filenames_both_processed(self, client):
        html = post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "same.docx"),
            (io.BytesIO(make_docx_bytes([WEAK_RESUME])), "same.docx"),
        ]).data.decode()
        assert "Ranked Candidates (2)" in html
        assert "same (1).docx" in html

    def test_oversize_file_rejected(self, client):
        big = b"%PDF-1.4\n" + b"0" * (11 * 1024 * 1024)
        resp = client.post("/screen",
                           data={"job_description": SAMPLE_JD,
                                 "resumes": [(io.BytesIO(big), "huge.pdf")]},
                           content_type="multipart/form-data",
                           follow_redirects=True)
        assert "No resumes could be processed" in resp.data.decode()

    def test_more_than_20_files_capped(self, client):
        files = [(io.BytesIO(make_docx_bytes(["Python Flask SQL %d" % i])),
                  "f%d.docx" % i) for i in range(21)]
        html = post_screen(client, SAMPLE_JD, files).data.decode()
        assert "Ranked Candidates (20)" in html
        assert "Maximum" in html
        assert app_mod.LAST_SUMMARY["total"] == 21

    def test_short_jd_warns_but_processes(self, client):
        html = post_screen(client, "Python developer role", [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "a.docx")
        ]).data.decode()
        assert "Ranked Candidates (1)" in html

    def test_resume_script_content_is_escaped(self, client):
        payload = ("<script>alert(2)</script> Python Flask SQL developer "
                   "with databases")
        html = post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([payload])), "xss.docx")
        ]).data.decode()
        assert "<script>alert(2)" not in html
        assert "&lt;script&gt;" in html

    def test_filename_angle_brackets_escaped(self, client):
        html = post_screen(client, SAMPLE_JD, [
            (io.BytesIO(b"junk"), "a<b.pdf"),
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "ok.docx"),
        ]).data.decode()
        assert "a&lt;b.pdf" in html
        assert "a<b.pdf" not in html

    def test_unicode_filename_and_content(self, client):
        html = post_screen(client, SAMPLE_JD + " développement", [
            (io.BytesIO(make_docx_bytes(["Python Flask développeur SQL"])),
             "résumé.docx"),
        ]).data.decode()
        assert "résumé.docx" in html

    def test_clear_resets_state(self, client):
        post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "a.docx")])
        assert len(app_mod.LAST_RESULTS) == 1
        resp = client.get("/clear", follow_redirects=True)
        assert resp.status_code == 200
        assert app_mod.LAST_RESULTS == []

    def test_screen_rejects_get(self, client):
        assert client.get("/screen").status_code == 405


# --------------------------------------------------------------------------
# 5. Overview dashboard statistics (latest screening, empty-safe)
# --------------------------------------------------------------------------

class TestOverviewStats:
    def _overview(self, client):
        return client.get("/").data.decode()

    def test_empty_state_before_any_screening(self, client):
        client.get("/clear")
        html = self._overview(client)
        assert html.count("<strong>—</strong>") >= 3
        assert "Awaiting first screening" in html

    def test_stats_reflect_first_batch(self, client):
        client.get("/clear")
        post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "a.docx"),
            (io.BytesIO(make_docx_bytes([WEAK_RESUME])), "b.docx"),
        ])
        html = self._overview(client)
        assert "<strong>2</strong>" in html  # uploaded + processed
        top = app_mod.LAST_SUMMARY["top_score"]
        assert ("%.2f%%" % top) in html
        assert "Awaiting first screening" not in html

    def test_stats_update_on_second_batch_not_accumulated(self, client):
        client.get("/clear")
        post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "a.docx"),
            (io.BytesIO(make_docx_bytes([WEAK_RESUME])), "b.docx"),
        ])
        first_top = app_mod.LAST_SUMMARY["top_score"]
        # Second batch: 3 uploaded, 1 invalid -> 2 processed.
        post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "c1.docx"),
            (io.BytesIO(make_docx_bytes([WEAK_RESUME])), "c2.docx"),
            (io.BytesIO(b"MZ junk"), "bad.exe"),
        ])
        html = self._overview(client)
        assert "<strong>3</strong>" in html  # uploaded follows latest batch
        assert "<strong>2</strong>" in html  # processed follows latest batch
        assert "<strong>5</strong>" not in html  # must not accumulate
        assert "<strong>4</strong>" not in html
        second_top = app_mod.LAST_SUMMARY["top_score"]
        assert ("%.2f%%" % second_top) in html
        if abs(first_top - second_top) > 0.001:
            assert ("%.2f%%" % first_top) not in html  # stale value gone

    def test_highest_matches_ranked_results(self, client):
        client.get("/clear")
        post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "a.docx"),
            (io.BytesIO(make_docx_bytes([WEAK_RESUME])), "b.docx"),
        ])
        expected = max(r["score"] for r in app_mod.LAST_RESULTS)
        assert app_mod.LAST_SUMMARY["top_score"] == expected
        assert ("%.2f%%" % expected) in self._overview(client)

    def test_revisits_do_not_double_count(self, client):
        client.get("/clear")
        post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "a.docx"),
        ])
        first = self._overview(client)
        assert self._overview(client) == first
        assert self._overview(client) == first

    def test_clear_restores_empty_state(self, client):
        post_screen(client, SAMPLE_JD, [
            (io.BytesIO(make_docx_bytes([STRONG_RESUME])), "a.docx"),
        ])
        assert "Awaiting first screening" not in self._overview(client)
        client.get("/clear")
        html = self._overview(client)
        assert html.count("<strong>—</strong>") >= 3
        assert "Awaiting first screening" in html
